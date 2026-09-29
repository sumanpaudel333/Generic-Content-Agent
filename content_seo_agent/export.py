"""
Approved product descriptions as a spreadsheet: SKU, product name, description.

The description is the one that goes to the product page, not a re-wording of
it. A published row uses the exact HTML that was sent to Odoo; an approved row
that has not been sent yet is assembled the same way publishing would assemble
it, delivery and disclaimer lines included.

Two ways to write the description:

    text   readable in Excel -- paragraphs, headings and "•" bullets
    html   exactly as sent to Odoo, for pasting or importing elsewhere

Spreadsheet programs run a cell that starts with "=" as a formula, and product
names and model-written text are not trusted input. The Excel file stores every
cell as text, so nothing is ever evaluated. CSV has no such setting, so a CSV
cell that could be read as a formula gets a leading apostrophe, the usual guard.
"""
import csv
import html
import io
import re
from datetime import datetime

from content_seo_agent import review_queue
from content_seo_agent.assembler import assemble_html
from content_seo_agent.constants import Status

FORMAT_XLSX = "xlsx"
FORMAT_CSV = "csv"
FORMATS = {FORMAT_XLSX: "Excel (.xlsx)", FORMAT_CSV: "CSV (.csv)"}

SCOPE_ALL = "all"
SCOPE_PUBLISHED = "published"
SCOPE_UNPUBLISHED = "unpublished"
SCOPES = {SCOPE_ALL: "All approved", SCOPE_PUBLISHED: "Published to Odoo",
          SCOPE_UNPUBLISHED: "Approved, not yet published"}

TEXT_PLAIN = "text"
TEXT_HTML = "html"
TEXTS = {TEXT_PLAIN: "Plain text", TEXT_HTML: "HTML, as sent to Odoo"}

HEADERS = ("SKU", "Product name", "Description")

MEDIA_TYPES = {
    FORMAT_XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    FORMAT_CSV: "text/csv; charset=utf-8",
}

# Excel refuses to open a file with a longer cell.
EXCEL_CELL_LIMIT = 32767
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_NUMBER = re.compile(r"^[+-]?\d+(\.\d+)?$")
_WRITTEN_FIELDS = ("overview", "features", "applications", "sections")


def approved_rows(scope: str = SCOPE_ALL) -> list[dict]:
    """Approved rows for the scope, sorted by product name."""
    published = {SCOPE_PUBLISHED: True, SCOPE_UNPUBLISHED: False}.get(scope)
    rows = review_queue.list_rows(status=Status.APPROVED, published=published)
    rows.sort(key=lambda r: ((r.get("title") or "").lower(), str(r.get("product_id") or "")))
    return rows


def description_html(row: dict) -> str:
    """The description as it is, or would be, on the product page. "" when the
    row has nothing written (a classification-only row, say)."""
    if row.get("published") and row.get("assembled_html"):
        return row["assembled_html"]
    draft = row.get("parsed_output") or {}
    if not isinstance(draft, dict) or not any(draft.get(f) for f in _WRITTEN_FIELDS):
        return ""
    return assemble_html(title=row.get("title") or "", draft_json=draft,
                         is_regulated=bool(row.get("is_regulated")),
                         ratio=row.get("ratio") or "",
                         quantity_detail=row.get("quantity_detail") or "")


def html_to_text(markup: str) -> str:
    """Readable text from description HTML: paragraphs separated by a blank
    line, each heading directly above its "•" bullets."""
    text = re.sub(r">\s+<", "><", markup or "")
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    # A heading ("Features & Benefits:") sits directly above its bullets; an
    # ordinary paragraph followed by a list keeps the blank line.
    text = re.sub(r"(?i):((?:</\w+>)*)</p>(?=<(?:ul|ol)\b)", r":\1\n", text)
    text = re.sub(r"(?i)<li\b[^>]*>", "• ", text)
    text = re.sub(r"(?i)</li>", "\n", text)
    text = re.sub(r"(?i)</(p|ul|ol|div|h[1-6])>", "\n\n", text)
    text = html.unescape(re.sub(r"(?s)<[^>]+>", "", text))
    lines = []
    for line in text.splitlines():
        line = re.sub(r"[ \t ]+", " ", line).strip()
        if line or (lines and lines[-1]):
            lines.append(line)
    return "\n".join(lines).strip()


def _csv_safe(value: str) -> str:
    if value.startswith(_FORMULA_START) and not _NUMBER.match(value):
        return "'" + value
    return value


def _to_csv(records: list[tuple]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(HEADERS)
    for record in records:
        writer.writerow([_csv_safe(v) for v in record])
    # The byte-order mark is what makes Excel read the file as UTF-8, so "•",
    # "m³" and accented names do not turn into gibberish when it is opened.
    return buffer.getvalue().encode("utf-8-sig")


def _to_xlsx(records: list[tuple]) -> tuple[bytes, int]:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Approved descriptions"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="004495")
        cell.alignment = Alignment(vertical="center")

    truncated = 0
    wrap = Alignment(wrap_text=True, vertical="top")
    for r, record in enumerate(records, start=2):
        for c, value in enumerate(record, start=1):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
            if len(value) > EXCEL_CELL_LIMIT:
                value = value[:EXCEL_CELL_LIMIT - 1] + "…"
                truncated += 1
            cell = sheet.cell(row=r, column=c)
            cell.value = value
            cell.data_type = "s"          # text, never a formula
            cell.alignment = wrap

    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 100
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:C{max(1, len(records) + 1)}"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), truncated


def build(scope: str = SCOPE_ALL, fmt: str = FORMAT_XLSX, text: str = TEXT_PLAIN,
          now: datetime | None = None) -> dict:
    """The file to download. Raises ValueError for an option that does not exist."""
    if scope not in SCOPES or fmt not in FORMATS or text not in TEXTS:
        raise ValueError("Unknown export option.")
    records, skipped = [], 0
    for row in approved_rows(scope):
        markup = description_html(row)
        if not markup:
            skipped += 1
            continue
        description = markup if text == TEXT_HTML else html_to_text(markup)
        records.append((str(row.get("product_id") or ""), row.get("title") or "", description))

    truncated = 0
    if fmt == FORMAT_XLSX:
        content, truncated = _to_xlsx(records)
    else:
        content = _to_csv(records)
    stamp = (now or datetime.now()).strftime("%Y-%m-%d")
    suffix = {SCOPE_ALL: "", SCOPE_PUBLISHED: "-published", SCOPE_UNPUBLISHED: "-not-published"}[scope]
    return {"content": content, "media_type": MEDIA_TYPES[fmt],
            "filename": f"approved-descriptions{suffix}-{stamp}.{fmt}",
            "count": len(records), "skipped": skipped, "truncated": truncated}
